"""Read author workbooks without modifying them."""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / ".tools/xls-reader"))
import openpyxl
import xlrd

UPSTREAM = ROOT / "data/manylabs1/upstream/7a1334c053201ccdd2054417f74aa3d62776dd03"

book = openpyxl.load_workbook(UPSTREAM / "MLforR.xlsx", read_only=True, data_only=True)
print("INPUT SHEETS", book.sheetnames)
for sheet in [book.worksheets[0]]:
    rows = list(sheet.iter_rows(values_only=True))
    print("INPUT", sheet.title, "rows", len(rows), "columns", sheet.max_column)
    for index, row in enumerate(rows):
        print(index + 1, json.dumps(row, ensure_ascii=True, default=str))
book.close()
reference = xlrd.open_workbook(str(UPSTREAM / "ML_ESCIfromR.xls"))
print("REFERENCE SHEETS", reference.sheet_names())
sheet = reference.sheet_by_index(0)
print("REFERENCE", sheet.name, sheet.nrows, sheet.ncols)
for index in range(min(sheet.nrows, 43)):
    print(index + 1, json.dumps(sheet.row_values(index), ensure_ascii=True))
