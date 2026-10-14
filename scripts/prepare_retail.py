"""Read-only workbook extraction to typed SQLite, with row-level decisions."""
import hashlib
import json
import sqlite3
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DIRECTORY = ROOT / "data/retail/uci-352-v1"


def main():
    manifest = json.loads((DIRECTORY/"source_manifest.json").read_bytes())
    source = DIRECTORY/"Online Retail.xlsx"
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if digest != manifest["files"][source.name]["sha256"]: raise ValueError("Workbook source drift")
    target = DIRECTORY/"transactions.sqlite"
    if target.exists(): raise ValueError("Prepared snapshot exists; do not overwrite")
    book = load_workbook(source, read_only=True, data_only=True)
    sheet = book.active
    records = sheet.iter_rows(values_only=True)
    headers = next(records)
    expected = ("InvoiceNo","StockCode","Description","Quantity","InvoiceDate","UnitPrice","CustomerID","Country")
    if tuple(headers) != expected: raise ValueError("Unexpected source headers")
    db = sqlite3.connect(target)
    db.execute('CREATE TABLE transactions(source_row INTEGER PRIMARY KEY, month TEXT, invoice TEXT, product TEXT, country TEXT, qty INTEGER, price_mills INTEGER, amount_mills INTEGER, kind TEXT, reason TEXT, customer_missing INTEGER, duplicate_candidate INTEGER)')
    count, issues, seen, decimals, monthly = 0, Counter(), set(), Counter(), Counter()
    first, last = None, None
    batch = []
    for source_row, row in enumerate(records, 2):
        invoice, product, description, qty, date, price, customer, country = row
        count += 1
        invoice = "" if invoice is None else str(invoice).strip()
        product = "" if product is None else str(product).strip()
        country = "" if country is None else str(country).strip()
        reason = None
        if not isinstance(date, datetime) or not invoice or not product or not country: reason = "missing_required"
        if type(qty) not in (int,float) or qty != int(qty): reason = "invalid_quantity"
        try:
            p = Decimal(str(price))
            if not p.is_finite(): raise ValueError("nonfinite")
            scale = max(0,-p.as_tuple().exponent)
            decimals[scale] += 1
            if p*1000 != (p*1000).to_integral_value(): reason = "price_precision_over_3"
            mills = int(p*1000)
        except Exception:
            mills = 0
            reason = "invalid_price"
        quantity = int(qty) if type(qty) in (int,float) and qty == int(qty) else 0
        if reason is None:
            if mills <= 0: reason = "nonpositive_price"
            elif quantity == 0: reason = "zero_quantity"
            elif invoice.upper().startswith("C") and quantity > 0: reason = "cancellation_sign_conflict"
            elif not invoice.upper().startswith("C") and quantity < 0: reason = "negative_without_cancellation"
        kind = "excluded" if reason else "cancellation" if invoice.upper().startswith("C") else "sale"
        month = date.strftime("%Y-%m") if isinstance(date,datetime) else ""
        if isinstance(date,datetime):
            first = date if first is None else min(first,date)
            last = date if last is None else max(last,date)
        # Candidate exact duplicates are retained, not silently deleted.
        signature = (invoice,product,description,qty,date,price,customer,country)
        duplicate = int(signature in seen)
        seen.add(signature)
        if duplicate: issues["duplicate_candidate_retained"] += 1
        if customer is None: issues["missing_customer_retained_or_classified"] += 1
        if description is None: issues["missing_description"] += 1
        if reason: issues[reason] += 1
        issues["kind_"+kind] += 1
        monthly[month] += 1
        batch.append((source_row,month,invoice,product,country,quantity,mills,quantity*mills,kind,reason,int(customer is None),duplicate))
        if len(batch) >= 10000:
            db.executemany('INSERT INTO transactions VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',batch)
            batch.clear()
        if count % 100000 == 0: print(f"Parsed {count} rows",flush=True)
    if batch: db.executemany('INSERT INTO transactions VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',batch)
    db.commit()
    db.close()
    book.close()
    if hashlib.sha256(source.read_bytes()).hexdigest() != digest: raise ValueError("Source changed during parse")
    record = {"source_sha256":digest,"prepared_sha256":hashlib.sha256(target.read_bytes()).hexdigest(),"sheet":sheet.title,"source_range":f"A1:H{count+1}","rows":count,"first_timestamp":first.isoformat(),"last_timestamp":last.isoformat(),"price_decimal_places":dict(decimals),"quality_counts":dict(issues),"monthly_source_rows":dict(sorted(monthly.items())),"currency":"GBP","exact_amount_unit":"0.001 GBP; no per-line cent rounding","duplicates":"exact raw-row candidates retained, not proven duplicate events","customer_policy":"missing CustomerID does not exclude sale; identifiers not copied into prepared DB","status":"prepared","created_at":datetime.now(timezone.utc).isoformat()}
    (DIRECTORY/"prepared_manifest.json").write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(record,ensure_ascii=True))


if __name__ == "__main__": main()
